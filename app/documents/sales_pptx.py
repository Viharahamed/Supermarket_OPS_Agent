# app/documents/sales_pptx.py

"""PPTX Sales Analysis Presentation Generator using python-pptx and Matplotlib charts.

Builds an 8-slide executive sales analysis PowerPoint deck from deterministic
reporting metrics and database entities.
"""

from __future__ import annotations

from datetime import datetime, date, timedelta
from decimal import Decimal
import os
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.database import get_db_context
from app.db.models import OwnerPreference, Store
from app.documents.charts import (
    create_daily_sales_chart,
    create_payment_method_chart,
    create_top_products_chart,
)
from app.exceptions import (
    DocumentGenerationError,
    InvalidDateError,
    InvalidDateRangeError,
)
from app.services import reporting_service, inventory_service


def _get_reports_dirs() -> tuple[Path, Path]:
    """Return platform-independent Path objects for reports and temporary chart files."""
    settings = get_settings()
    base_dir = Path(settings.local_document_dir)
    reports_dir = base_dir / "reports"
    temp_charts_dir = reports_dir / "temp_charts"
    reports_dir.mkdir(parents=True, exist_ok=True)
    temp_charts_dir.mkdir(parents=True, exist_ok=True)
    return reports_dir, temp_charts_dir


def parse_date(date_str: str) -> date:
    """Parse YYYY-MM-DD date string safely."""
    try:
        return datetime.strptime(date_str, "%Y-%m-%d").date()
    except Exception as exc:
        raise InvalidDateError(f"Invalid date format '{date_str}'. Expected YYYY-MM-DD.") from exc


def generate_sales_analysis_pptx(
    start_date_str: Optional[str] = None,
    end_date_str: Optional[str] = None,
    days: int = 7,
    db: Session = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> dict:
    """Generate an 8-slide PowerPoint Sales Analysis presentation."""
    start_date_str = start_date_str or start_date
    end_date_str = end_date_str or end_date
    settings = get_settings()

    try:
        from zoneinfo import ZoneInfo
        tz = ZoneInfo(settings.timezone)
        today = datetime.now(tz).date()
    except Exception:
        today = date.today()

    if end_date_str:
        end_d = parse_date(end_date_str)
    else:
        end_d = today

    if start_date_str:
        start_d = parse_date(start_date_str)
    else:
        start_d = end_d - timedelta(days=days - 1)

    if start_d > end_d:
        raise InvalidDateRangeError(f"Start date ({start_d}) cannot be after end date ({end_d}).")

    if db is None:
        with get_db_context() as session:
            return _generate_pptx_impl(session, start_d, end_d)
    return _generate_pptx_impl(db, start_d, end_d)


def _generate_pptx_impl(db: Session, start_d: date, end_d: date) -> dict:
    # 1. Retrieve Store Profile
    store = db.query(Store).filter(Store.id == 1).first()
    store_name = store.name if store else "Lakshmi Kirana & General Store"

    # 2. Retrieve Deterministic Reporting Data
    # inclusive start, exclusive end (end_d + 1 day)
    exclusive_end_d = end_d + timedelta(days=1)

    summary_report = reporting_service.get_sales_summary(start_d, exclusive_end_d)
    payment_report = reporting_service.get_payment_breakdown(start_d, exclusive_end_d)
    top_products_list = reporting_service.get_top_products(start_d, exclusive_end_d, limit=5, by="revenue")
    low_stock_items = inventory_service.get_low_stock(db)

    # 3. Generate Daily Sales Trend Data Points
    daily_sales_data = []
    curr = start_d
    while curr <= end_d:
        try:
            d_rep = reporting_service.get_daily_sales(curr)
            daily_sales_data.append({
                "date": curr.strftime("%d %b"),
                "total_sales": float(d_rep.grand_total),
                "bill_count": d_rep.bill_count,
            })
        except Exception:
            daily_sales_data.append({
                "date": curr.strftime("%d %b"),
                "total_sales": 0.0,
                "bill_count": 0,
            })
        curr += timedelta(days=1)

    # 4. Render Matplotlib Charts
    import io
    from app.storage import get_storage

    reports_dir, temp_charts_dir = _get_reports_dirs()
    daily_chart_path = create_daily_sales_chart(daily_sales_data, temp_charts_dir / "daily_sales.png")
    
    payment_dict = {
        "CASH": float(payment_report.cash),
        "UPI": float(payment_report.upi),
        "CARD": float(payment_report.card),
        "KHATA": float(payment_report.khata),
    }
    payment_chart_path = create_payment_method_chart(payment_dict, temp_charts_dir / "payment_methods.png")

    top_prod_dict_list = [
        {"name": item.product_name, "total_revenue": float(item.sales_value)}
        for item in top_products_list
    ]
    top_prod_chart_path = create_top_products_chart(top_prod_dict_list, temp_charts_dir / "top_products.png")

    # 5. Build PowerPoint Deck using python-pptx
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9 Widescreen aspect ratio

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text: str, subtitle_text: str = ""):
        # Header background banner
        banner = slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(10), Inches(0.8)  # MSO_SHAPE.RECTANGLE = 1
        )
        banner.fill.solid()
        banner.fill.fore_color.rgb = RGBColor(26, 54, 93)  # #1A365D Deep Blue
        banner.line.color.rgb = RGBColor(26, 54, 93)

        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(9), Inches(0.6))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(11)
            p2.font.color.rgb = RGBColor(226, 232, 240)

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(5.625))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(26, 54, 93)
    bg1.line.color.rgb = RGBColor(26, 54, 93)

    tbox1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(8.0), Inches(2.5))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "WEEKLY SALES & BUSINESS ANALYSIS"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)

    p1_sub = tf1.add_paragraph()
    p1_sub.text = f"Store: {store_name}\nPeriod: {start_d.strftime('%d %b %Y')} to {end_d.strftime('%d %b %Y')}"
    p1_sub.font.size = Pt(16)
    p1_sub.font.color.rgb = RGBColor(226, 232, 240)

    # -------------------------------------------------------------------------
    # SLIDE 2: Executive Summary
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Executive Sales Summary", f"{start_d.strftime('%d %b')} - {end_d.strftime('%d %b %Y')}")

    # Metrics Box Table Layout
    metrics = [
        ("Total Sales", f"₹ {summary_report.total_sales:,.2f}"),
        ("Finalized Bills", f"{summary_report.total_bills}"),
        ("Avg Bill Value", f"₹ {summary_report.average_bill_value:,.2f}"),
        ("Taxable Sales", f"₹ {summary_report.total_taxable_amount:,.2f}"),
        ("Total GST Tax", f"₹ {summary_report.total_tax:,.2f}"),
    ]

    lefts = [0.8, 3.8, 6.8, 2.3, 5.3]
    tops = [1.3, 1.3, 1.3, 3.3, 3.3]
    widths = [2.6, 2.6, 2.6, 2.6, 2.6]
    heights = [1.6, 1.6, 1.6, 1.6, 1.6]

    for i, (label, val_str) in enumerate(metrics):
        card = slide2.shapes.add_shape(1, Inches(lefts[i]), Inches(tops[i]), Inches(widths[i]), Inches(heights[i]))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(247, 250, 252)
        card.line.color.rgb = RGBColor(43, 108, 176)
        card.line.width = Pt(1.5)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_c1 = tf_c.paragraphs[0]
        p_c1.text = label.upper()
        p_c1.font.size = Pt(10)
        p_c1.font.bold = True
        p_c1.font.color.rgb = RGBColor(113, 128, 150)
        p_c1.alignment = PP_ALIGN.CENTER

        p_c2 = tf_c.add_paragraph()
        p_c2.text = val_str
        p_c2.font.size = Pt(18)
        p_c2.font.bold = True
        p_c2.font.color.rgb = RGBColor(43, 108, 176)
        p_c2.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------------------
    # SLIDE 3: Daily Sales Trend
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Daily Sales Performance Trend", "Track daily revenue trajectory")

    slide3.shapes.add_picture(str(daily_chart_path), Inches(0.5), Inches(1.1), Inches(5.2), Inches(3.8))

    txt3 = slide3.shapes.add_textbox(Inches(5.9), Inches(1.3), Inches(3.8), Inches(3.5))
    tf3 = txt3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "Key Observations:"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(26, 54, 93)

    bullets3 = [
        f"Period Total Revenue: ₹ {summary_report.total_sales:,.2f}",
        f"Total Finalized Orders: {summary_report.total_bills}",
        f"Highest Order Value: ₹ {summary_report.highest_bill_value:,.2f}",
        f"Lowest Order Value: ₹ {summary_report.lowest_bill_value:,.2f}",
    ]
    for b in bullets3:
        p_b = tf3.add_paragraph()
        p_b.text = f"• {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = RGBColor(45, 55, 72)

    # -------------------------------------------------------------------------
    # SLIDE 4: Payment Method Breakdown
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Payment Collection Channels", "Breakdown by Cash, UPI, Card, and Khata Credit")

    slide4.shapes.add_picture(str(payment_chart_path), Inches(0.5), Inches(1.1), Inches(5.2), Inches(3.8))

    txt4 = slide4.shapes.add_textbox(Inches(5.9), Inches(1.3), Inches(3.8), Inches(3.5))
    tf4 = txt4.text_frame
    tf4.word_wrap = True
    p4 = tf4.paragraphs[0]
    p4.text = "Channel Breakdown:"
    p4.font.size = Pt(14)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(26, 54, 93)

    payment_bullets = [
        f"CASH Sales: ₹ {payment_report.cash:,.2f}",
        f"UPI Digital Sales: ₹ {payment_report.upi:,.2f}",
        f"CARD Sales: ₹ {payment_report.card:,.2f}",
        f"KHATA Credit Sales: ₹ {payment_report.khata:,.2f}",
    ]
    for pb in payment_bullets:
        p_pb = tf4.add_paragraph()
        p_pb.text = f"• {pb}"
        p_pb.font.size = Pt(12)
        p_pb.font.color.rgb = RGBColor(45, 55, 72)

    # -------------------------------------------------------------------------
    # SLIDE 5: Top Products by Revenue
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Top Selling Products", "Highest revenue contributors during period")

    slide5.shapes.add_picture(str(top_prod_chart_path), Inches(0.5), Inches(1.1), Inches(5.2), Inches(3.8))

    txt5 = slide5.shapes.add_textbox(Inches(5.9), Inches(1.3), Inches(3.8), Inches(3.5))
    tf5 = txt5.text_frame
    tf5.word_wrap = True
    p5 = tf5.paragraphs[0]
    p5.text = "Top Revenue Items:"
    p5.font.size = Pt(14)
    p5.font.bold = True
    p5.font.color.rgb = RGBColor(26, 54, 93)

    if top_products_list:
        for tp in top_products_list[:4]:
            p_tp = tf5.add_paragraph()
            p_tp.text = f"• {tp.product_name}: {tp.quantity_sold:.0f} sold (₹{tp.sales_value:,.2f})"
            p_tp.font.size = Pt(11)
            p_tp.font.color.rgb = RGBColor(45, 55, 72)
    else:
        p_tp = tf5.add_paragraph()
        p_tp.text = "• No product sales recorded in period."
        p_tp.font.size = Pt(12)

    # -------------------------------------------------------------------------
    # SLIDE 6: GST Tax Collection Summary
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "GST Tax Collection Breakdown", "Deterministic CGST and SGST tax metrics")

    gst_table_shape = slide6.shapes.add_table(5, 2, Inches(1.5), Inches(1.4), Inches(7.0), Inches(2.8))
    gst_table = gst_table_shape.table
    gst_table.columns[0].width = Inches(4.5)
    gst_table.columns[1].width = Inches(2.5)

    gst_rows = [
        ("Tax Metric", "Amount (₹)"),
        ("Total Net Taxable Amount", f"₹ {summary_report.total_taxable_amount:,.2f}"),
        ("Central GST (CGST)", f"₹ {summary_report.total_cgst:,.2f}"),
        ("State GST (SGST)", f"₹ {summary_report.total_sgst:,.2f}"),
        ("Total GST Tax Collected", f"₹ {summary_report.total_tax:,.2f}"),
    ]
    for r_idx, (col1, col2) in enumerate(gst_rows):
        cell1 = gst_table.cell(r_idx, 0)
        cell2 = gst_table.cell(r_idx, 1)

        cell1.text = col1
        cell2.text = col2

        if r_idx == 0:
            cell1.fill.solid()
            cell1.fill.fore_color.rgb = RGBColor(43, 108, 176)
            cell2.fill.solid()
            cell2.fill.fore_color.rgb = RGBColor(43, 108, 176)

    # -------------------------------------------------------------------------
    # SLIDE 7: Inventory Health & Reorder Alerts
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Inventory Health & Reorder Alerts", "Products at or below reorder threshold")

    txt7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(8.4), Inches(4.0))
    tf7 = txt7.text_frame
    tf7.word_wrap = True
    p7 = tf7.paragraphs[0]
    p7.text = f"Low Stock Items Identified: {len(low_stock_items)}"
    p7.font.size = Pt(14)
    p7.font.bold = True
    p7.font.color.rgb = RGBColor(197, 48, 48)  # Red warning

    if low_stock_items:
        for ls in low_stock_items[:6]:
            p_ls = tf7.add_paragraph()
            p_ls.text = f"• {ls.product_name} (SKU: {ls.sku}) — Current Stock: {ls.stock_quantity:.2f} {ls.unit} (Reorder Level: {ls.reorder_level:.2f})"
            p_ls.font.size = Pt(11)
            p_ls.font.color.rgb = RGBColor(45, 55, 72)
    else:
        p_ls = tf7.add_paragraph()
        p_ls.text = "✅ All catalog inventory levels are healthy above reorder thresholds."
        p_ls.font.size = Pt(12)
        p_ls.font.color.rgb = RGBColor(56, 161, 105)

    # -------------------------------------------------------------------------
    # SLIDE 8: Deterministic Business Insights
    # -------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Store Business Insights", "Actionable operational recommendations")

    txt8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(8.4), Inches(4.0))
    tf8 = txt8.text_frame
    tf8.word_wrap = True
    p8 = tf8.paragraphs[0]
    p8.text = "Operational Insights & Recommendations:"
    p8.font.size = Pt(14)
    p8.font.bold = True
    p8.font.color.rgb = RGBColor(26, 54, 93)

    top_p_name = top_products_list[0].product_name if top_products_list else "Top category items"
    insights = [
        f"Sales Volume: Generated ₹{summary_report.total_sales:,.2f} across {summary_report.total_bills} finalized orders (Avg: ₹{summary_report.average_bill_value:,.2f}).",
        f"Revenue Driver: '{top_p_name}' emerged as the leading revenue contributor during this period.",
        f"Inventory Action: {len(low_stock_items)} items require immediate restocking to prevent stock-outs.",
        f"Tax Compliance: Total GST collected ₹{summary_report.total_tax:,.2f} (CGST: ₹{summary_report.total_cgst:,.2f}, SGST: ₹{summary_report.total_sgst:,.2f}).",
    ]
    for ins in insights:
        p_ins = tf8.add_paragraph()
        p_ins.text = f"• {ins}"
        p_ins.font.size = Pt(12)
        p_ins.font.color.rgb = RGBColor(45, 55, 72)

    # 6. Save PPTX to BytesIO buffer
    pptx_buffer = io.BytesIO()
    prs.save(pptx_buffer)
    pptx_bytes = pptx_buffer.getvalue()

    # 7. Store binary presentation using DocumentStorage abstraction
    file_name = f"sales_analysis_{start_d.isoformat()}_{end_d.isoformat()}.pptx"
    relative_path = f"reports/{file_name}"

    try:
        storage = get_storage()
        result = storage.save(
            content=pptx_bytes,
            relative_path=relative_path,
            artifact_type="sales_pptx",
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        return {
            "success": True,
            "file_path": result.file_path,
            "file_name": result.file_name,
            "relative_path": result.relative_path,
            "start_date": start_d.isoformat(),
            "end_date": end_d.isoformat(),
            "total_sales": f"{summary_report.total_sales:.2f}",
            "total_bills": summary_report.total_bills,
            "slides_count": len(prs.slides),
        }
    except Exception as exc:
        raise DocumentGenerationError(f"Failed to generate PPTX sales report: {exc}") from exc
