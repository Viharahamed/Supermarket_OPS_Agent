# tests/test_documents.py
"""Tests for Phase 13B.1 Document Generation Hardening & Telegram Delivery.

Tests PDF invoice generation, PDF magic header validation, financial correctness,
fractional quantity formatting, PPTX sales presentation slide inspection,
empty-data presentation handling, tool registry integration, and Telegram document delivery.
"""

from decimal import Decimal
import os
import pytest
from unittest.mock import AsyncMock, MagicMock, patch
from pptx import Presentation

from app.db.database import get_db_context
from app.db.models import Product, Bill, BillItem, Customer
from app.documents.invoice_pdf import generate_invoice_pdf, _sanitize_filename
from app.documents.sales_pptx import generate_sales_analysis_pptx
from app.exceptions import BillNotFoundError, BillNotFinalizedError
from app.services.billing_service import create_draft_bill, add_bill_item, finalize_bill
from app.services.inventory_service import create_product
from app.tools import registry


@pytest.fixture
def setup_finalized_bill(db_session):
    """Fixture to create products, customer, draft bill, and finalized bill with mixed GST rates."""
    p1 = create_product(
        name="Tata Salt 1kg",
        sku="SALT-TATA-1KG",
        category="Grocery",
        cost_price=Decimal("20.00"),
        mrp=Decimal("28.00"),
        selling_price=Decimal("25.00"),
        stock_quantity=Decimal("100"),
        hsn_code="25010010",
        gst_rate=Decimal("5.00"),
    )
    p2 = create_product(
        name="Maggi 2-Min Noodle",
        sku="MAGGI-70G",
        category="Snacks",
        cost_price=Decimal("11.00"),
        mrp=Decimal("14.00"),
        selling_price=Decimal("14.00"),
        stock_quantity=Decimal("50"),
        hsn_code="19023010",
        gst_rate=Decimal("12.00"),
    )

    with get_db_context() as db:
        cust = Customer(name="Ramesh Sharma", phone="9876543210")
        db.add(cust)
        db.commit()
        db.refresh(cust)
        cust_id = cust.id

    bill = create_draft_bill(customer_id=cust_id)
    bill = add_bill_item(bill.id, p1.id, Decimal("2"))
    bill = add_bill_item(bill.id, p2.id, Decimal("3"))
    final_bill = finalize_bill(bill.id, payment_method="CASH")
    return final_bill, bill.id


# -----------------------------------------------------------------------------
# PDF Invoice Tests
# -----------------------------------------------------------------------------

def test_generate_invoice_pdf_success(setup_finalized_bill):
    """Test successful generation of PDF invoice for a finalized bill."""
    final_bill, bill_id = setup_finalized_bill

    res = generate_invoice_pdf(bill_id)
    assert res["bill_id"] == bill_id
    assert "file_path" in res
    assert os.path.exists(res["file_path"])
    assert res["file_name"].endswith(".pdf")
    assert os.path.getsize(res["file_path"]) > 1000  # Non-trivial PDF size

    # Verify PDF magic header bytes (%PDF-)
    with open(res["file_path"], "rb") as f:
        header = f.read(4)
        assert header == b"%PDF"


def test_generate_invoice_pdf_financial_correctness(setup_finalized_bill):
    """Verify generated PDF invoice returns exact grand total matching DB entity."""
    final_bill, bill_id = setup_finalized_bill

    res = generate_invoice_pdf(bill_id)
    assert Decimal(res["grand_total"]) == final_bill.grand_total
    assert res["payment_method"] == "CASH"


def test_generate_invoice_pdf_fractional_quantity(db_session):
    """Verify PDF invoice generation with fractional item quantity (e.g. 0.5 kg)."""
    p = create_product(
        name="Loose Sugar 1kg",
        sku="SUGAR-LOOSE-1KG",
        category="Grocery",
        cost_price=Decimal("35.00"),
        mrp=Decimal("45.00"),
        selling_price=Decimal("40.00"),
        stock_quantity=Decimal("50.00"),
        unit="kg",
        is_loose=True,
    )
    draft = create_draft_bill()
    draft = add_bill_item(draft.id, p.id, Decimal("0.50"))
    finalized = finalize_bill(draft.id, payment_method="UPI")

    res = generate_invoice_pdf(finalized.id)
    assert os.path.exists(res["file_path"])
    assert res["grand_total"] == f"{finalized.grand_total:.2f}"


def test_generate_invoice_pdf_draft_fails(db_session):
    """Test that generating PDF for a draft bill raises BillNotFinalizedError."""
    p = create_product(name="Sugar 1kg", sku="SUGAR-1KG", selling_price=Decimal("45.00"), mrp=Decimal("50.00"))
    draft_bill = create_draft_bill()
    draft_bill = add_bill_item(draft_bill.id, p.id, Decimal("1"))

    with pytest.raises(BillNotFinalizedError):
        generate_invoice_pdf(draft_bill.id)


def test_generate_invoice_pdf_not_found(db_session):
    """Test that generating PDF for a non-existent bill raises BillNotFoundError."""
    with pytest.raises(BillNotFoundError):
        generate_invoice_pdf(99999)


def test_sanitize_filename_utility():
    """Verify filename sanitization removes dangerous/illegal characters."""
    unsafe = "BILL/2026:09*07?<1001>|win"
    clean = _sanitize_filename(unsafe)
    assert "/" not in clean
    assert ":" not in clean
    assert "*" not in clean
    assert "?" not in clean
    assert "<" not in clean
    assert ">" not in clean
    assert "|" not in clean
    assert clean == "BILL202609071001win"


# -----------------------------------------------------------------------------
# PPTX Sales Presentation Tests
# -----------------------------------------------------------------------------

def test_generate_sales_analysis_pptx_success(setup_finalized_bill):
    """Test successful generation of an 8-slide PPTX sales analysis presentation."""
    res = generate_sales_analysis_pptx(start_date=None, end_date=None)

    assert "file_path" in res
    file_path = res["file_path"]
    assert os.path.exists(file_path)
    assert res["file_name"].endswith(".pptx")
    assert os.path.getsize(file_path) > 5000

    # Verify slide count & structure using python-pptx
    prs = Presentation(file_path)
    assert len(prs.slides) == 8
    assert res["slides_count"] == 8

    # Inspect slide titles
    slide_titles = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                slide_titles.append(shape.text_frame.text)
                break

    assert any("sales" in t.lower() and "analysis" in t.lower() for t in slide_titles)
    assert any("executive" in t.lower() for t in slide_titles)


def test_generate_sales_analysis_pptx_empty_data(db_session):
    """Test PPTX generation when there are no sales records."""
    res = generate_sales_analysis_pptx(start_date="2020-01-01", end_date="2020-01-02")
    assert os.path.exists(res["file_path"])
    prs = Presentation(res["file_path"])
    assert len(prs.slides) == 8


# -----------------------------------------------------------------------------
# Tool Registry Execution Tests
# -----------------------------------------------------------------------------

def test_tool_registry_generate_invoice_pdf(setup_finalized_bill):
    """Test executing generate_invoice_pdf tool via ToolRegistry."""
    final_bill, bill_id = setup_finalized_bill

    tool_res = registry.execute("generate_invoice_pdf", {"bill_id": bill_id})
    assert tool_res.success is True
    assert "file_path" in tool_res.data
    assert os.path.exists(tool_res.data["file_path"])


def test_tool_registry_generate_sales_pptx(setup_finalized_bill):
    """Test executing generate_sales_analysis_pptx tool via ToolRegistry."""
    tool_res = registry.execute("generate_sales_analysis_pptx", {})
    assert tool_res.success is True
    assert "file_path" in tool_res.data
    assert os.path.exists(tool_res.data["file_path"])


# -----------------------------------------------------------------------------
# Telegram Document Delivery Handler Tests
# -----------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_telegram_handler_delivers_document_attachment(tmp_path):
    """Test that Telegram text_message_handler sends document attachment when tool generates a file."""
    from app.telegram.handlers import text_message_handler
    from app.agent.schemas import AgentResponse

    # Create dummy PDF file
    pdf_path = tmp_path / "invoice_bill_1.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 dummy pdf content")

    mock_agent_response = AgentResponse(
        content="Here is your requested PDF invoice for Bill #1.",
        metadata={
            "tool_results": [
                {
                    "tool_name": "generate_invoice_pdf",
                    "data": {
                        "file_path": str(pdf_path),
                        "file_name": "invoice_bill_1.pdf",
                    },
                }
            ]
        },
    )

    update = MagicMock()
    update.message.text = "Generate PDF invoice for bill 1"
    update.effective_user.id = 12345
    update.effective_chat.id = 67890
    update.message.reply_text = AsyncMock()
    update.message.reply_document = AsyncMock()

    context = MagicMock()

    with patch("app.telegram.handlers.handle_message", return_value=mock_agent_response):
        await text_message_handler(update, context)

    # Assert reply_text sent message
    update.message.reply_text.assert_called()
    # Assert reply_document sent file attachment
    update.message.reply_document.assert_called_once()
    kwargs = update.message.reply_document.call_args.kwargs
    assert kwargs["filename"] == "invoice_bill_1.pdf"
